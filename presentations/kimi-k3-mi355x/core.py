from pathlib import Path
import json, math, copy, re, textwrap
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

ROOT=Path(__file__).resolve().parent
OUT=ROOT/"output"
OUT.mkdir(exist_ok=True)
SLIDES=[]
for f in sorted(ROOT.glob("slides-*.json")):
    SLIDES.extend(json.loads(f.read_text()))
SOURCES=json.loads((ROOT/"sources.json").read_text())
REPO="ajith-sirra-amd/InferenceMAX_Rocm_Team"
BRANCH="docs/kimi-k3-mi355x-teaching-20261029"
BASE=f"https://github.com/{REPO}/blob/{BRANCH}/presentations/kimi-k3-mi355x"
CHAPTERS={"00":"Orientation","01":"Attention fundamentals","02":"Fused attention",
"03":"Kimi-K3 architecture","04":"MI355X execution","05":"Distributed serving",
"06":"FP8 data flow","07":"Accuracy and execution proof","08":"Performance reasoning",
"09":"Putting it together","10":"Reference"}
C={"bg":"0B1220","card":"152239","card2":"1C2B44","line":"33465F","white":"F2F6FC",
"muted":"A7B8CD","fp8":"51D79B","bf16":"69AEFF","fp32":"B799FF",
"amber":"F4BF61","teal":"52CFD1","neutral":"CBD7E6","red":"FA818E"}
FONT="DejaVu Sans"
W,H=13.333333,7.5
def rgb(c): return RGBColor.from_string(C.get(c,c))
def el(tag,**attrs):
    e=etree.Element("{http://schemas.openxmlformats.org/presentationml/2006/main}"+tag)
    for k,v in attrs.items():e.set(k,str(v))
    return e
def sub(parent,tag,**attrs):
    e=el(tag,**attrs);parent.append(e);return e
class Canvas:
    def __init__(self,prs,data,index):
        self.slide=prs.slides.add_slide(prs.slide_layouts[6]);self.d=data;self.i=index
        self.groups=[];self.geometry=[]
        self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=rgb("bg")
        self.rect(0,0,W,.075,"teal")
        self.text(.5,.23,11,.28,CHAPTERS[data["chapter"]].upper(),10,"teal",bold=True)
        title_size=29
        title_font_path="/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
        while ImageFont.truetype(title_font_path,title_size).getlength(data["title"])>12.0*72*.94 and title_size>19:
            title_size-=1
        title_shape=self.text(.5,.69,12.3,.65,data["title"],title_size,"white",bold=True)
        title_shape.text_frame.word_wrap=False
        self.text(.52,1.47,12.1,.43,data["subtitle"],15,"muted")
        self.rect(.5,7.06,12.33,.012,"line")
        self.text(.5,7.15,4.2,.18,"KIMI-K3  /  MI355X  •  29 OCT 2026",8,"muted")
        guide=self.text(6.8,7.12,4.8,.25,"READ CHAPTER  ↗",9,"teal",align=PP_ALIGN.RIGHT)
        guide.text_frame.paragraphs[0].runs[0].hyperlink.address=self.guide_url()
        self.text(11.9,7.12,.9,.25,f"{index:02d} / {len(SLIDES)}",9,"muted",align=PP_ALIGN.RIGHT)
        if data.get("formula"):
            self.rect(.5,6.36,12.33,.48,"card2")
            self.text(.7,6.45,11.9,.28,data["formula"],13,"white",bold=True)
        else:
            self.text(.6,6.52,11.5,.22,"CLICK TO BUILD THE DIAGRAM  •  REFERENCES AND EXPLANATION IN SPEAKER NOTES",8,"muted")
    def guide_url(self):
        return f'{BASE}/guide/{self.d["chapter"]}-{slug(CHAPTERS[self.d["chapter"]])}.md#s-{self.i:02d}'
    def rect(self,x,y,w,h,fill="card",border=None,round=False):
        shp=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,
            Inches(x),Inches(y),Inches(w),Inches(h))
        shp.fill.solid();shp.fill.fore_color.rgb=rgb(fill)
        if border:shp.line.color.rgb=rgb(border);shp.line.width=Pt(1.15)
        else:shp.line.fill.background()
        if round:
            try:shp.adjustments[0]=.09
            except:pass
        self.geometry.append((shp.shape_id,x,y,w,h))
        return shp
    def circle(self,x,y,d,fill="teal"):
        shp=self.slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
        shp.fill.solid();shp.fill.fore_color.rgb=rgb(fill);shp.line.fill.background()
        self.geometry.append((shp.shape_id,x,y,d,d));return shp
    def text(self,x,y,w,h,text,size=18,color="white",bold=False,align=PP_ALIGN.LEFT):
        shp=self.slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=shp.text_frame;tf.clear();tf.word_wrap=True
        tf.margin_left=tf.margin_right=Inches(.015);tf.margin_top=tf.margin_bottom=0
        for k,line in enumerate(str(text).split("\n")):
            p=tf.paragraphs[0] if k==0 else tf.add_paragraph()
            p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(0)
            r=p.add_run();r.text=line;r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
        self.geometry.append((shp.shape_id,x,y,w,h));return shp
    def arrow(self,x1,y1,x2,y2,color="muted"):
        ln=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        ln.line.color.rgb=rgb(color);ln.line.width=Pt(1.8)
        line=ln._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
        end=etree.SubElement(line,"{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
        end.set("type","triangle");end.set("w","sm");end.set("len","sm")
        return ln
    def reveal(self,*shapes): self.groups.append([x.shape_id for x in shapes if x is not None])
    def card(self,x,y,w,h,title,body="",color="teal",big=False):
        a=self.rect(x,y,w,h,"card",color,True)
        b=self.rect(x+.17,y+.2,.06,.35,color)
        title_size=30 if big else 21
        available=(w-.5)*72*.93
        font_path="/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
        while title_size>12:
            font=ImageFont.truetype(font_path,title_size);lines=1;line=""
            too_wide=False
            for word in title.split():
                if font.getlength(word)>available:too_wide=True
                candidate=(line+" "+word).strip()
                if line and font.getlength(candidate)>available:lines+=1;line=word
                else:line=candidate
            if not too_wide and lines*title_size*1.2 <= (68 if big else 49):break
            title_size-=1
        c=self.text(x+.34,y+.2,w-.5,.72 if not big else .95,title,title_size,color,True)
        d=self.text(x+.24,y+1.0,w-.46,h-1.12,body,16,"white") if body else None
        return [a,b,c,d]
    def label(self,x,y,w,text,color="teal"):
        a=self.rect(x,y,w,.34,"card2",color,True)
        b=self.text(x+.06,y+.065,w-.12,.2,text,9,color,True,PP_ALIGN.CENTER)
        return [a,b]
    def finish(self):
        note=self.d["note"]+"\n\nANIMATION: Advance one click at a time to reveal each group. In the video the same groups appear progressively.\n\nREAD MORE: "+self.guide_url()+"\n\nSOURCES:\n"
        note+="\n".join(SOURCES[k]["label"]+": "+SOURCES[k]["url"] for k in self.d.get("refs",[]))
        self.slide.notes_slide.notes_text_frame.text=note
        return self
def slug(s):return re.sub(r"[^a-z0-9]+","-",s.lower()).strip("-")
def make_prs():
    p=Presentation();p.slide_width=Inches(W);p.slide_height=Inches(H)
    p.core_properties.title="Kimi-K3 on AMD MI355X — Architecture and FP8 Execution"
    p.core_properties.subject="Conceptual teaching deck with linked technical handbook"
    p.core_properties.author="Ajith Sirra"
    p.core_properties.keywords="Kimi-K3, MI355X, KDA, MLA, FP8, FMHA, ROCm, vLLM"
    for part in p.part.package.iter_parts():
        if str(part.partname).startswith("/ppt/theme/"):
            doc=etree.fromstring(part.blob)
            ns={"a":"http://schemas.openxmlformats.org/drawingml/2006/main"}
            for tag in ["hlink","folHlink"]:
                for node in doc.xpath("//a:clrScheme/a:"+tag+"/*",namespaces=ns):
                    node.tag="{"+ns["a"]+"}srgbClr";node.attrib.clear();node.set("val",C["teal"])
            part._blob=etree.tostring(doc,xml_declaration=True,encoding="UTF-8",standalone=True)
    return p
def add_animation(slide,groups):
    # Standard DrawingML entrance effects, grouped into explicit mouse-click builds.
    root=slide._element
    trans=el("transition",spd="med",advClick="1");sub(trans,"fade");root.append(trans)
    timing=el("timing");tn=sub(timing,"tnLst");par=sub(tn,"par")
    main=sub(par,"cTn",id="1",dur="indefinite",restart="never",nodeType="tmRoot")
    mainchildren=sub(main,"childTnLst");seq=sub(mainchildren,"seq",concurrent="1",nextAc="seek")
    seqtime=sub(seq,"cTn",id="2",dur="indefinite",nodeType="mainSeq")
    effects=sub(seqtime,"childTnLst");nid=3
    for group in groups:
        outer=sub(effects,"par");ot=sub(outer,"cTn",id=nid,fill="hold");nid+=1
        conds=sub(ot,"stCondLst");sub(conds,"cond",delay="indefinite")
        middle=sub(sub(ot,"childTnLst"),"par");mt=sub(middle,"cTn",id=nid,fill="hold");nid+=1
        sub(sub(mt,"stCondLst"),"cond",delay="0")
        children=sub(mt,"childTnLst")
        for spid in group:
            ep=sub(children,"par")
            et=sub(ep,"cTn",id=nid,presetID="10",presetClass="entr",presetSubtype="0",fill="hold",
                   grpId="0",nodeType="clickEffect");nid+=1
            sub(sub(et,"stCondLst"),"cond",delay="0")
            acts=sub(et,"childTnLst")
            st=sub(acts,"set");bh=sub(st,"cBhvr");ct=sub(bh,"cTn",id=nid,dur="1",fill="hold");nid+=1
            sub(sub(ct,"stCondLst"),"cond",delay="0")
            sub(sub(bh,"tgtEl"),"spTgt",spid=spid)
            an=sub(bh,"attrNameLst");sub(an,"attrName").text="style.visibility"
            sub(sub(st,"to"),"strVal",val="visible")
            ae=sub(acts,"animEffect",transition="in",filter="fade")
            abh=sub(ae,"cBhvr");sub(abh,"cTn",id=nid,dur="400");nid+=1
            sub(sub(abh,"tgtEl"),"spTgt",spid=spid)
    for tag,evt in [("prevCondLst","onPrev"),("nextCondLst","onNext")]:
        cond=sub(sub(seq,tag),"cond",evt=evt,delay="0");sub(sub(cond,"tgtEl"),"sldTgt")
    root.append(timing)
def duplicate_visible(source,prs,hidden):
    dest=prs.slides.add_slide(prs.slide_layouts[6])
    dest.background.fill.solid();dest.background.fill.fore_color.rgb=rgb("bg")
    for shape in source.shapes:
        if shape.shape_id in hidden:continue
        e=copy.deepcopy(shape.element)
        # Staging deck needs no external hyperlinks/relationships.
        for h in e.xpath(".//a:hlinkClick"):
            h.getparent().remove(h)
        dest.shapes._spTree.insert_element_before(e,"p:extLst")
    return dest
